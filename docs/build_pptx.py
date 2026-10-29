"""Build the WhyCode prototype PPT (hackFront India 2026, FORTEC).

9 slides, 16:9, from the verified facts of the Django graph build.
Run with the pptx venv: ~/workspace/.venvs/pptx/bin/python docs/build_pptx.py
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

OUT = "/home/hatch/workspace/goals/hackathon-registration-drive/projects/hackfront-fortec/docs/whycode-pitch-deck.pptx"
DARK = RGBColor(0x0B, 0x1F, 0x1A)
GREEN = RGBColor(0x4A, 0xDE, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY = RGBColor(0xD1, 0xD5, 0xDB)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def slide(title, bullets, subtitle=None):
    s = prs.slides.add_slide(blank)
    bg = s.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = DARK
    tx = s.shapes.add_textbox(Inches(0.6), Inches(0.35),
                              Inches(12.1), Inches(1.0)).text_frame
    tx.word_wrap = True
    p = tx.paragraphs[0]
    p.text = title
    p.font.size = Pt(34)
    p.font.bold = True
    p.font.color.rgb = WHITE
    y = 1.5
    if subtitle:
        tx2 = s.shapes.add_textbox(Inches(0.6), Inches(1.15),
                                   Inches(12.1), Inches(0.6)).text_frame
        tx2.word_wrap = True
        p = tx2.paragraphs[0]
        p.text = subtitle
        p.font.size = Pt(20)
        p.font.color.rgb = GREEN
        y = 1.9
    box = s.shapes.add_textbox(Inches(0.6), Inches(y),
                               Inches(12.1), Inches(5.2)).text_frame
    box.word_wrap = True
    for i, b in enumerate(bullets):
        p = box.paragraphs[0] if i == 0 else box.add_paragraph()
        p.text = "•  " + b
        p.font.size = Pt(20)
        p.font.color.rgb = GREY
        p.space_after = Pt(14)
    return s


slide("WhyCode", [
    "Answers one question every developer asks: why is this code like this?",
    "Team FORTEC — hackFront India 2026, Theme: AI and Developer Tools (T-1).",
    "Ronit Shrimankar (lead), Anshul Kamthe, Udita Singh.",
], subtitle="Built on the real history of django/django")

slide("The problem", [
    "Codebases accumulate decisions nobody can explain any more.",
    "A new developer reads an odd function and guesses why it exists.",
    "A wrong guess breaks something the original author fixed on purpose.",
    "The reasons live scattered across commits, pull requests and tickets —",
    "never in one place, never linked to the code they explain.",
])

slide("The solution: WhyCode", [
    "Ingest a public repo's commits, PRs and tickets into a graph.",
    "Code entities linked to the decisions, people and incidents behind them.",
    "Ask a plain-language 'why' question and get a cited evidence chain,",
    "a confidence score — and a refusal when the graph holds no evidence.",
    "No invented answers, ever. The refusal is a feature, not a gap.",
])

slide("How it works", [
    "1. Ingest — GitHub API via gh CLI (commits, PRs) plus Trac tickets,",
    "   because Django tracks decisions in Trac, not GitHub Issues.",
    "2. Graph — entities (commit, PR, ticket, file, person) and edges",
    "   (authored, fixes, references, touches, supersedes).",
    "3. Ask — match the question to entities, walk the evidence chain,",
    "   flag decisions superseded by newer ones, refuse when evidence is absent.",
])

slide("Demo: real answers from the Django graph", [
    "'Why does ignore_warnings match substrings by default?' — 0.99,",
    "  chained to Trac ticket #37089 and its fixing commit.",
    "'Why is there a transitional option to disable multiline template tags?'",
    "  — 0.86, chained to ticket #35899.",
    "'Why is Python 3.15 compatibility being added?' — 0.99,",
    "  25 linked records around ticket #36664.",
])

slide("The refusal, honestly demonstrated", [
    "'Why does Django use a microservices architecture?' — REFUSED.",
    "The term appears in no ingested record, so the engine refuses",
    "instead of answering from the nearest generic match.",
    "An earlier version answered this from an unrelated docs commit;",
    "the absent-term rule plus a regression test fixed that for good.",
])

slide("Engineering: runnable, tested, no magic", [
    "Python standard library only; SQLite graph; no paid APIs, no accounts.",
    "The bundled database runs the whole demo with zero credentials.",
    "10/10 automated tests, including refusal and supersession cases.",
    "Graph today: 200 commits, 100 PRs, 100 Trac tickets, 161 files,",
    "90 people, 635+ edges — all from Django's real, public history.",
])

slide("Who it helps, and what is next", [
    "Onboarding developers, reviewers and maintainers stop guessing at",
    "history and start citing it.",
    "Next: function-level entities (today: file-level), Django Enhancement",
    "Proposals as design-doc evidence, and a Neo4j adapter for the",
    "partner track. The local graph already runs without any of that.",
])

slide("WhyCode — by FORTEC", [
    "Repo: github.com/eres0045/hackfront-fortec (public, with setup steps).",
    "Demo: python -m whycode.cli --db data/whycode.db --question 'Why ...?'",
    "Built during the hackFront league timeline, from Django's real history.",
    "Thank you — questions welcome.",
])

prs.save(OUT)
print(f"saved {OUT} with {len(list(prs.slides))} slides")
